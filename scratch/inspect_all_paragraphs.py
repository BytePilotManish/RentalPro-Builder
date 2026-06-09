import docx

doc = docx.Document(r"d:\rental pro\new\19-05-2026 - kan (2.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P {idx}: {text[:150]}")
