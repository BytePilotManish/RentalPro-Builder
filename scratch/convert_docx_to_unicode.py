import sys
import os
import re
import docx

# Import mapping from knconverter
mapping = {
    "C"     : "ಅ", "D"     : "ಆ", "E"     : "ಇ", "F"     : "ಈ", "G"     : "ಉ", "H"     : "ಊ",
    "IÄ"    : "ಋ", "J"     : "ಎ", "K"     : "ಏ", "L"     : "ಐ", "M"     : "ಒ", "N"     : "ಓ",
    "O"     : "ಔ", "A"     : "ಂ", "B"     : "ಃ", "Pï"    : "ಕ್", "PÀ"    : "ಕ", "PÁ"    : "ಕಾ", 
    "Q"     : "ಕಿ", "PÉ"    : "ಕೆ", "PË"    : "ಕೌ", "Sï"    : "ಖ್", "R"     : "ಖ", "SÁ"    : "ಖಾ",
    "T"     : "ಖಿ", "SÉ"    : "ಖೆ", "SË"    : "ಖೌ", "Uï"    : "ಗ್", "UÀ"    : "ಗ", "UÁ"    : "ಗಾ",
    "V"     : "ಗಿ", "UÉ"    : "ಗೆ", "UË"    : "ಗೌ", "Wï"    : "ಘ್", "WÀ"    : "ಘ", "WÁ"    : "ಘಾ",
    "X"     : "ಘಿ", "WÉ"    : "ಘೆ", "WË"    : "ಘೌ", "k"     : "ಞ", "Zï"    : "ಚ್", "ZÀ"    : "ಚ",
    "ZÁ"    : "ಚಾ", "a"     : "ಚಿ", "ZÉ"    : "ಚೆ", "ZË"    : "ಚೌ", "bï"    : "ಛ್", "bÀ"    : "ಛ",
    "bÁ"    : "ಛಾ", "c"     : "ಛಿ", "bÉ"    : "ಛೆ", "bË"    : "ಛೌ", "eï"    : "ಜ್", "d"     : "ಜ",
    "eÁ"    : "ಜಾ", "f"     : "ಜಿ", "eÉ"    : "ಜೆ", "eË"    : "ಜೌ", "gÀhiï" : "ಝ್", "gÀhÄ"  : "ಝ",
    "gÀhiÁ" : "ಝಾ", "jhÄ"   : "ಝಿ", "gÉhÄ"  : "ಝೆ", "gÉhÆ"  : "ಝೊ", "gÀhiË" : "ಝೌ", "Y"     : "ಙ",
    "mï"    : "ಟ್", "l"     : "ಟ", "mÁ"    : "ಟಾ", "n"     : "ಟಿ", "mÉ"    : "ಟೆ", "mË"    : "ಟೌ",
    "oï"    : "ಠ್", "oÀ"    : "ಠ", "oÁ"    : "ಠಾ", "p"     : "ಠಿ", "oÉ"    : "ಠೆ", "oË"    : "ಠೌ",
    "qï"    : "ಡ್", "qÀ"    : "ಡ", "qÁ"    : "ಡಾ", "r"     : "ಡಿ", "qÉ"    : "ಡೆ", "qË"    : "ಡೌ",
    "qsï"   : "ಢ್", "qsÀ"   : "ಢ", "qsÁ"   : "ಢಾ", "rü"    : "ಢಿ", "qsÉ"   : "ಢೆ", "qsË"   : "ಢೌ",
    "uï"    : "ಣ್", "t"     : "ಣ", "uÁ"    : "ಣಾ", "tÂ"    : "ಣಿ", "uÉ"    : "ಣೆ", "uË"    : "ಣೌ",
    "vï"    : "ತ್", "vÀ"    : "ತ", "vÁ"    : "ತಾ", "w"     : "ತಿ", "vÉ"    : "ತೆ", "vË"    : "ತೌ",
    "xï"    : "ಥ್", "xÀ"    : "ಥ", "xÁ"    : "ಥಾ", "y"     : "ಥಿ", "xÉ"    : "ಥೆ", "xË"    : "ಥೌ",
    "zï"    : "ದ್", "zÀ"    : "ದ", "zÁ"    : "ದಾ", "¢"     : "ದಿ", "zÉ"    : "ದೆ", "zË"    : "ದೌ",
    "zsï"   : "ಧ್", "zsÀ"   : "ಧ", "zsÁ"   : "ಧಾ", "¢ü"    : "ಧಿ", "zsÉ"   : "ಧೆ", "zsË"   : "ಧೌ",
    "£ï"    : "ನ್", "£À"    : "ನ", "£Á"    : "ನಾ", "¤"     : "ನಿ", "£É"    : "ನೆ", "£Ë"    : "ನೌ",
    "¥ï"    : "ಪ್", "¥À"    : "ಪ", "¥Á"    : "ಪಾ", "¦"     : "ಪಿ", "¥É"    : "ಪೆ", "¥Ë"    : "ಪೌ",
    "¥sï"   : "ಫ್", "¥sÀ"   : "ಫ", "¥sÁ"   : "ಫಾ", "¦ü"    : "ಫಿ", "¥sÉ"   : "ಫೆ", "¥sË"   : "ಫೌ",
    "¨ï"    : "ಬ್", "§"     : "ಬ", "¨Á"    : "ಬಾ", "©"     : "ಬಿ", "¨É"    : "ಬೆ", "¨Ë"    : "ಬೌ",
    "¨sï"   : "ಭ್", "¨sÀ"   : "ಭ", "¨sÁ"   : "ಭಾ", "©ü"    : "ಭಿ", "¨sÉ"   : "ಭೆ", "¨sË"   : "ಭೌ",
    "ªÀiï"  : "ಮ್", "ªÀÄ"   : "ಮ", "ªÀiÁ"  : "ಮಾ", "«Ä"    : "ಮಿ", "ªÉÄ"   : "ಮೆ", "ªÀiË"  : "ಮೌ",
    "AiÀiï" : "ಯ್", "AiÀÄ"  : "ಯ", "0iÀÄ"  : "ಯ", "AiÀiÁ" : "ಯಾ", "0iÀiÁ" : "ಯಾ", "¬Ä"    : "ಯಿ",
    "0iÀÄÄ" : "ಯು", "AiÉÄ"  : "ಯೆ", "0iÉÆ"  : "ಯೊ", "AiÉÆ"  : "ಯೊ", "AiÀiË" : "ಯೌ", "gï"    : "ರ್",
    "gÀ"    : "ರ", "gÁ"    : "ರಾ", "j"     : "ರಿ", "gÉ"    : "ರೆ", "gË"    : "ರೌ", "¯ï"    : "ಲ್",
    "®"     : "ಲ", "¯Á"    : "ಲಾ", "°"     : "ಲಿ", "¯É"    : "ಲೆ", "¯Ë"    : "ಲೌ", "ªï"    : "ವ್",
    "ªÀ"    : "ವ", "ªÁ"    : "ವಾ", "«"     : "ವಿ", "ªÀÅ"   : "ವು", "ªÀÇ"   : "ವೂ", "ªÉ"    : "ವೆ",
    "ªÉÃ"   : "ವೇ", "ªÉÊ"   : "ವೈ", "ªÉÆ"   : "ಮೊ", "ªÉÆÃ"  : "ಮೋ", "ªÉÇ"   : "ವೊ", "ªÉÇÃ"  : "ವೋ",
    "ªÉ  "  : "ವೆ", "¥ÀÅ"   : "ಪು", "¥ÀÇ"   : "ಪೂ", "¥sÀÅ"  : "ಫು", "¥sÀÇ"  : "ಫೂ", "ªË"    : "ವೌ",
    "±ï"    : "ಶ್", "±À"    : "ಶ", "±Á"    : "ಶಾ", "²"     : "ಶಿ", "±É"    : "ಶೆ", "±Ë"    : "ಶೌ",
    "µï"    : "ಷ್", "µÀ"    : "ಷ", "µÁ"    : "ಷಾ", "¶"     : "ಷಿ", "µÉ"    : "ಷೆ", "µË"    : "ಷೌ",
    "¸ï"    : "ಸ್", "¸À"    : "ಸ", "¸Á"    : "ಸಾ", "¹"     : "ಸಿ", "¸É"    : "ಸೆ", "¸Ë"    : "ಸೌ",
    "ºï"    : "ಹ್", "ºÀ"    : "ಹ", "ºÁ"    : "ಹಾ", "»"     : "ಹಿ", "ºÉ"    : "ಹೆ", "ºË"    : "ಹೌ",
    "¼ï"    : "ಳ್", "¼À"    : "ಳ", "¼Á"    : "ಳಾ", "½"     : "ಳಿ", "¼É"    : "ಳೆ", "¼Ë"    : "ಳೌ"
}

broken_cases = {
    "Ã": {"value": "ೀ", "mapping": {"ಿ": "ೀ", "ೆ": "ೇ", "ೊ": "ೋ"}},
    "Ä": {"value": "ು", "mapping": {}},
    "Æ": {"value": "ೂ", "mapping": {"ೆ": "ೊ"}},
    "È": {"value": "ೃ", "mapping": {}},
    "Ê": {"value": "ೈ", "mapping": {"ೆ": "ೈ"}}
}

dependent_vowels = ["್", "ಾ", "ಿ", "ೀ", "ು", "ೂ", "ೃ", "ೆ", "ೇ", "ೈ", "ೊ", "ೋ", "ೌ"]
ignore_list = {"ö": "", "÷": ""}

vattaksharagalu = {
    "Ì" : "ಕ", "Í" : "ಖ", "Î" : "ಗ", "Ï" : "ಘ", "Õ" : "ಞ", "Ñ" : "ಚ", "Ò" : "ಛ", "Ó" : "ಜ",
    "Ô" : "ಝ", "Ö" : "ಟ", "×" : "ಠ", "Ø" : "ಡ", "Ù" : "ಢ", "Ú" : "ಣ", "Û" : "ತ", "Ü" : "ಥ",
    "Ý" : "ದ", "Þ" : "ಧ", "ß" : "ನ", "à" : "ಪ", "á" : "ಫ", "â" : "ಬ", "ã" : "ಭ", "ä" : "ಮ",
    "å" : "ಯ", "æ" : "ರ", "è" : "ಲ", "é" : "ವ", "ê" : "ಶ", "ë" : "ಷ", "ì" : "ಸ", "í" : "ಹ",
    "î" : "ಳ", "ç" : "ರ"
}

ascii_arkavattu = {"ð": "ರ"}

def process_vattakshara(letters, t):
    last_letter = letters[-1] if len(letters) > 0 else ""
    if last_letter in dependent_vowels:    
        letters[-1] = "್"
        letters.append(vattaksharagalu[t])
        letters.append(last_letter)
    else:
        letters.append("್")
        letters.append(vattaksharagalu[t])
    return letters

def process_arkavattu(letters, t):
    last_letter = letters[-1] if len(letters) > 0 else ""
    second_last = letters[-2] if len(letters) > 1 else ""
    if last_letter in dependent_vowels:    
        letters[-2] = ascii_arkavattu[t]
        letters[-1] = "್"
        letters.append(second_last)
        letters.append(last_letter)
    else:
        letters[-1] = ascii_arkavattu[t]
        letters.append("್")
        letters.append(last_letter)
    return letters

def process_broken_cases(letters, t):
    last_letter = letters[-1] if len(letters) > 0 else ""
    broken_case_mapping = broken_cases[t]["mapping"]
    if last_letter in broken_case_mapping:
        letters[-1] = broken_case_mapping[last_letter]
    else:
        letters.append(broken_cases[t]["value"])
    return letters

def find_mapping(op, txt, current_pos):
    max_len = 4
    remaining = len(txt) - current_pos
    if remaining < 5:
        max_len = (remaining - 1)
    
    n = 0
    for i in range(max_len, -1, -1):
        substr_till = current_pos + i + 1
        t = txt[current_pos:substr_till]
        
        if t in mapping:
            if len(op) > 0 and re.search("್$", op[-1]) is not None:
                op.append("‍")  # ZWJ
            op.append(mapping[t])
            n = i
            break
        else:
            if i > 0:
                continue
            
            op = list(''.join(op))
            if t in ascii_arkavattu:
                op = process_arkavattu(op, t)
            elif t in vattaksharagalu:
                op = process_vattakshara(op, t)
            elif t in broken_cases:
                op = process_broken_cases(op, t)
            else:
                op.append(t)
    return [n, op]

def process_word(word):
    i = 0
    max_len = len(word)
    op = []
    while i < max_len:
        if word[i] in ignore_list:
            i += 1
            continue
        data = find_mapping(op, word, i)
        op = data[1]
        i += (1 + data[0])
    return ''.join(op)

def process_line(line):
    line = line.strip()
    words = line.split(' ')
    op_words = []
    for word in words:
        op_words.append(process_word(word))
    return ' '.join(op_words)

# DOCX conversion main
def convert_docx_nudi_to_unicode(in_path, out_path):
    print(f"Loading {in_path}...")
    doc = docx.Document(in_path)
    
    print("Converting text to Unicode...")
    # Process only paragraphs 26 to 132 for the first agreement (0-indexed 25 to 131)
    new_doc = docx.Document()
    
    # Configure default style of target document
    style = new_doc.styles['Normal']
    style.font.name = 'Tunga'  # Unicode Kannada font
    style.font.size = docx.shared.Pt(12)
    
    p_count = 0
    for idx in range(25, 132):
        if idx >= len(doc.paragraphs):
            break
        p = doc.paragraphs[idx]
        new_p = new_doc.add_paragraph()
        p_count += 1
        
        # Copy paragraph formatting
        new_p.paragraph_format.alignment = p.paragraph_format.alignment
        new_p.paragraph_format.space_before = p.paragraph_format.space_before
        new_p.paragraph_format.space_after = p.paragraph_format.space_after
        new_p.paragraph_format.line_spacing = p.paragraph_format.line_spacing
        
        # Process runs in paragraph
        for r in p.runs:
            # Only convert if font is Nudi Akshar-02 or if it contains CP1252 Nudi chars
            # But actually, let's convert everything except numbers and standard English if possible,
            # or simply process all runs because CP1252 conversion won't affect standard plain English/numbers much
            # except that characters like letters might get mapped. So only convert runs with Nudi font or non-ascii
            # or if it's CP1252 text. Since the entire document is Kannada Nudi, we convert all runs.
            original_text = r.text
            if not original_text.strip():
                new_run = new_p.add_run(original_text)
                new_run.bold = r.bold
                new_run.italic = r.italic
                continue
                
            unicode_text = process_line(original_text)
            
            # Print a sample of conversion (safe)
            if p_count <= 5 and r.text.strip():
                safe_unicode = unicode_text.encode('ascii', 'backslashreplace').decode('ascii')
                print(f"  Sample Run: {repr(original_text)} -> {safe_unicode}")
                
            new_run = new_p.add_run(unicode_text)
            new_run.bold = r.bold
            new_run.italic = r.italic
            new_run.font.name = 'Tunga'  # Force Unicode Kannada font
            new_run.font.size = r.font.size or docx.shared.Pt(12)
            
    print(f"Processed {p_count} paragraphs. Saving to {out_path}...")
    new_doc.save(out_path)
    print("Done!")

if __name__ == "__main__":
    in_file = r"d:\rental pro\new\19-05-2026 - kan (2.docx"
    out_file = r"d:\rental pro\new\TEMPLATE_KAN_UNICODE.docx"
    convert_docx_nudi_to_unicode(in_file, out_file)
