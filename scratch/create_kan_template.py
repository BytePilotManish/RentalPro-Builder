import docx
import os
import sys

# Add root folder to sys.path to run from scratch
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

def make_template():
    in_file = r"d:\rental pro\new\19-05-2026 - kan (2.docx"
    out_file = r"d:\rental pro\new\TEMPLATE_KAN.docx"
    
    if not os.path.exists(in_file):
        print("Source file not found")
        return
        
    doc = docx.Document(in_file)
    new_doc = docx.Document()
    
    # Configure style
    style = new_doc.styles['Normal']
    style.font.name = 'Nudi Akshar-02'
    style.font.size = docx.shared.Pt(12)
    
    # We copy paragraphs 25 to 131 (which is index 25 to 131 inclusive)
    p_count = 0
    for idx in range(25, 132):
        if idx >= len(doc.paragraphs):
            break
        p = doc.paragraphs[idx]
        new_p = new_doc.add_paragraph()
        p_count += 1
        
        # Copy paragraph formatting
        new_p.paragraph_format.alignment = p.paragraph_format.alignment
        new_p.paragraph_format.space_before = p.paragraph_format.space_before
        new_p.paragraph_format.space_after = p.paragraph_format.space_after
        new_p.paragraph_format.line_spacing = p.paragraph_format.line_spacing
        
        # We will check if the paragraph text contains any strings to replace
        text = p.text
        
        # Replace hardcoded values with placeholders in Nudi ASCII representation
        # Paragraph 26: Date
        if idx == 26:
            # සನ್ ಎರಡು ಸಾವಿರದ ಇಪ್ಪತ್ತಾರನೇ ಇಸವಿ ಮೇ ಮಾಹೆ -> ಸನ್ {{AGREEMENT_YEAR_WORDS}} ಇಸವಿ {{AGREEMENT_MONTH}} ಮಾಹೆ
            # In Nudi: ¸À£ï JgÀqÀÄ ¸Á«gÀzÀ E¥ÀàvÁÛgÀ£ÉÃ E¸À« ªÉÄÃ ªÀiÁºÉ
            # JgÀqÀÄ ¸Á«gÀzÀ E¥ÀàvÁÛgÀ£ÉÃ -> {{AGREEMENT_YEAR_WORDS}}
            # ªÉÄÃ -> {{AGREEMENT_MONTH}}
            text = text.replace("JgÀqÀÄ ¸Á«gÀzÀ E¥ÀàvÁÛgÀ£ÉÃ", "{{AGREEMENT_YEAR_WORDS}}")
            text = text.replace("ªÉÄÃ", "{{AGREEMENT_MONTH}}")
            
        elif idx == 27:
            # ದಿನಾಂಕ ಹತ್ತೊಂಭತ್ತು (19-05-2026) ರಲ್ಲೂ -> ದಿನಾಂಕ {{AGREEMENT_DATE}} ರಲ್ಲೂ
            # In Nudi: ¢£ÁAPÀ ºÀvÉÆÛA¨sÀvÀÄÛ (19-05-2026) gÀ®Æè
            # ºÀvÉÆÛA¨sÀvÀÄÛ (19-05-2026) -> {{AGREEMENT_DATE}}
            text = text.replace("ºÀvÉÆÛA¨sÀvÀÄÛ (19-05-2026)", "{{AGREEMENT_DATE}}")
            
        elif idx == 28:
            # Owner details
            # Address: ¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51 £ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ
            # Parent: ¯ÉÃmï.ZÀAzÀæ¥Àà.©
            # Age: 41
            # Name: ²æÃ.µÀtÄäR¥Àà.©.¹
            text = text.replace("¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51 £ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ", "{{OWNER_ADDRESS}}")
            text = text.replace("¯ÉÃmï.ZÀAzÀæ¥Àà.©", "{{OWNER_PARENT}}")
            text = text.replace("41", "{{OWNER_AGE}}")
            text = text.replace("²æÃ.µÀtÄäR¥Àà.©.¹", "{{OWNER_NAME}}")
            
        elif idx == 43:
            # Tenant details
            # Address: ¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51 £ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ
            # Age: 52
            # Name: ²æÃªÀÄw.¥ÁªÀðvÀªÀÄä
            text = text.replace("¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51 £ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ", "{{TENANT_ADDRESS}}")
            text = text.replace("52", "{{TENANT_AGE}}")
            text = text.replace("²æÃªÀÄw.¥ÁªÀðvÀªÀÄä", "{{TENANT_NAME}}")
            
        elif idx == 45:
            # Property description and address
            # Address: ¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51£ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ
            # Description: £É® ªÀiÁºÀrAiÀÄ°ègÀÄªÀ Dgï.¹.¹ bÀvÀÄæ¥ÀÅöìåzÀ Dgï.¹.¹. bÀvÀÄæ... -> We can replace from "£É® ªÀiÁºÀrAiÀÄ°ègÀÄªÀ" to "C£ÀÄPÀÆ®«gÀÄªÀ"
            # In Nudi it is: £É® ªÀiÁºÀrAiÀÄ°ègÀÄªÀ Dgï.¹.¹ bÀvÀÄæ¥ÀÅöìåzÀ ¥ÀÇªÀðzÀ ¨ÁV®Ä¼Àî, MAz CrUÉ ªÀÄ£É, MAzÀÄ ºÁ¯ï, MAzÀÄ gÀÆªÀiï, ªÀÄvÀÄÛ ¨Ávï gÀÆªÀiï, ±ËZÁ®AiÀÄ, ºÁUÀÆ«zÀÄåvï «ÄÃlgi / ¤Ãj£À C£ÀÄPÀÆ®«gÀÄªÀ
            # Wait, let's just search and replace:
            text = text.replace("¨ÉAUÀ¼ÀÆgÀÄ – 562162, UÀAUÉÆAqÀ£ÀºÀ½î, ®PÀëöäöå¥ÀÄgÀ CAZÉ, ¯ÉÃPï gÀ¸ÉÛ, zsÀgÀägÁAiÀÄ gÀ¸ÉÛ, 51£ÉÃ £ÀA§gÀÄ¼Àî ²æÃ ¨sÀÆªÀgÀºÀ ¤®AiÀÄ", "{{PREMISES_ADDRESS}}")
            text = text.replace("£É® ªÀiÁºÀrAiÀÄ°ègÀÄªÀ Dgï.¹.¹ bÀvÀÄæ¥ÀÅöìåzÀ ¥ÀÇªÀðzÀ ¨ÁV®Ä¼Àî, MAzÀÄ CrUÉ ªÀÄ£É, MAzÀÄ ºÁ¯ï, MAzÀÄ gÀÆªÀiï, ªÀÄvÀÄÛ ¨Ávï gÀÆªÀiï, ±ËZÁ®AiÀÄ, ºÁUÀÆ«zÀÄåvï «ÄÃlgi / ¤Ãj£À C£ÀÄPÀÆ®«gÀÄªÀ", "{{PREMISES_DESCRIPTION}}")
            
        elif idx == 47:
            # Deposit
            # gÀ.20,000/- (E¥ÀàvÀÄÛ ¸Á«gÀ)
            text = text.replace("gÀ.20,000/- (E¥ÀàvÀÄÛ ¸Á«gÀ)", "gÀ.{{DEPOSIT_AMOUNT}}/- ({{DEPOSIT_AMOUNT_WORDS}})")
            
        elif idx == 48:
            # Rent amount & due day
            # gÀ.4,500/- (£Á®ÄÌ ¸Á«gÀzÀ L£ÀÆgÀ)
            # 28
            text = text.replace("gÀ.4,500/- (£Á®ÄÌ ¸Á«gÀzÀ L£ÀÆgÀ)", "gÀ.{{RENT_AMOUNT}}/- ({{RENT_AMOUNT_WORDS}})")
            text = text.replace("28", "{{RENT_PAYMENT_DAY}}")
            
        elif idx == 49:
            # Start date and duration
            # 15-01-2026
            # 11 (ºÀ£ÉÆßAzÀÄ)
            text = text.replace("15-01-2026", "{{LEASE_START_DATE}}")
            text = text.replace("11 (ºÀ£ÉÆßAzÀÄ)", "{{LEASE_PERIOD_NUM}} ({{LEASE_PERIOD}})")
            
        elif idx == 51:
            # Escalation rate
            # 5%
            text = text.replace("5%", "{{ESCALATION_RATE}}")
            
        elif idx == 54:
            # Notice period & lease duration in notice clause
            # 11
            # 3 (ªÀÄÆgÀÄ)
            text = text.replace("11", "{{LEASE_PERIOD_NUM}}")
            text = text.replace("3 (ªÀÄÆgÀÄ)", "{{NOTICE_PERIOD_NUM}} ({{NOTICE_PERIOD}})")
            
        elif idx == 60:
            # Tenant signature line
            text = text + "\n\n{{TENANT_SIG_NAMES}}"
            
        elif idx == 67:
            # Owner signature line
            text = text + "\n\n{{OWNER_SIG_NAMES}}"

        # Save runs to new paragraph
        if text.strip():
            run = new_p.add_run(text)
            run.font.name = 'Nudi Akshar-02'
            run.font.size = docx.shared.Pt(12)
            
    new_doc.save(out_file)
    print(f"Created template at {out_file} successfully.")

if __name__ == "__main__":
    make_template()
