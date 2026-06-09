import docx
import os
import sys

# Import mapping/functions from scratch.convert_docx_to_unicode
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from scratch.convert_docx_to_unicode import process_line

in_file = r"d:\rental pro\new\19-05-2026 - kan (2.docx"
if os.path.exists(in_file):
    doc = docx.Document(in_file)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    
    # Process from index 130 to end
    lines = []
    for idx in range(130, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        if p.text.strip():
            converted = process_line(p.text)
            lines.append(f"P {idx+1}: {converted}")
            
    with open("scratch/rest_kan_text_unicode.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("Written to scratch/rest_kan_text_unicode.txt")
else:
    print("File not found")
