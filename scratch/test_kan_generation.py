import docx
import os
import sys

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from app import replace_placeholders
from kannada_helper import unicode_to_nudi

def test_kan_gen():
    template_path = "TEMPLATE_KAN.docx"
    output_docx = "scratch/test_kan_output.docx"
    
    if not os.path.exists(template_path):
        print("TEMPLATE_KAN.docx not found!")
        return
        
    doc = docx.Document(template_path)
    
    unicode_payload = {
        "AGREEMENT_YEAR_WORDS": "ಎರಡು ಸಾವಿರದ ಇಪ್ಪತ್ತಾರನೇ",
        "AGREEMENT_MONTH": "ಮೇ",
        "AGREEMENT_DATE": "ಹತ್ತೊಂಭತ್ತು (19-05-2026)",
        "OWNER_NAME": "ಶ್ರೀ.ಷಣ್ಮುಖಪ್ಪ.ಬಿ.ಸಿ",
        "OWNER_PARENT": "ಲೇಟ್.ಚಂದ್ರಪ್ಪ.ಬಿ",
        "OWNER_AGE": "41",
        "OWNER_ADDRESS": "ಬೆಂಗಳೂರು–562162, ಗಂಗೊಂಡನಹಳ್ಳಿ, ಲಕ್ಷ್ಮಿಪುರ ಅಂಚೆ, ಲೇಕ್ ರಸ್ತೆ, ಧರ್ಮರಾಯ ರಸ್ತೆ, 51 ನೇ ನಂಬರುಳ್ಳ ಶ್ರೀ ಭೂವರಹ ನಿಲಯ",
        "TENANT_NAME": "ಶ್ರೀಮತಿ.ಪಾರ್ವತಮ್ಮ",
        "TENANT_PARENT": "ಗಂಡ ಲೇಟ್ ರಂಗಯ್ಯ",
        "TENANT_AGE": "52",
        "TENANT_ADDRESS": "ಬೆಂಗಳೂರು–562162, ಗಂಗೊಂಡನಹಳ್ಳಿ, ಲಕ್ಷ್ಮಿಪುರ ಅಂಚೆ, ಲೇಕ್ ರಸ್ತೆ, ಧರ್ಮರಾಯ ರಸ್ತೆ, 51 ನೇ ನಂಬರುಳ್ಳ ಶ್ರೀ ಭೂವರಹ ನಿಲಯ",
        "PREMISES_ADDRESS": "ಬೆಂಗಳೂರು–562162, ಗಂಗೊಂಡನಹಳ್ಳಿ, ಲಕ್ಷ್ಮಿಪುರ ಅಂಚೆ, ಲೇಕ್ ರಸ್ತೆ, ಧರ್ಮರಾಯ ರಸ್ತೆ, 51ನೇ ನಂಬರುಳ್ಳ ಶ್ರೀ ಭೂವರಹ ನಿಲಯ",
        "PREMISES_DESCRIPTION": "ನೆಲ ಮಾಹಡಿಯಲ್ಲಿರುವ ಆರ್.ಸಿ.ಸಿ ಛಾವಣಿಯುಳ್ಳ ಪೂರ್ವದ ಬಾಗಿಲುಳ್ಳ, ಒಂದು ಅಡಿಗೆ ಮನೆ, ಒಂದು ಹಾಲ್, ಒಂದು ರೂಮ್, ಮತ್ತು ಬಾತ್ ರೂಮ್, ಶೌಚಾಲಯ, ಹಾಗೂವಿದ್ಯುತ್ ಮೀಟರ್ / ನೀರಿನ ಅನುಕೂಲವಿರುವ",
        "DEPOSIT_AMOUNT": "20,000",
        "DEPOSIT_AMOUNT_WORDS": "ಇಪ್ಪತ್ತು ಸಾವಿರ",
        "RENT_AMOUNT": "4,500",
        "RENT_AMOUNT_WORDS": "ನಾಲ್ಕು ಸಾವಿರದ ಐನೂರು",
        "RENT_PAYMENT_DAY": "28",
        "LEASE_START_DATE": "15-01-2026",
        "LEASE_PERIOD_NUM": "11",
        "LEASE_PERIOD": "ಹನ್ನೊಂದು",
        "ESCALATION_RATE": "5%",
        "NOTICE_PERIOD_NUM": "3",
        "NOTICE_PERIOD": "ಮೂರು",
        "OWNER_SIG_NAMES": "ಷಣ್ಮುಖಪ್ಪ.ಬಿ.ಸಿ",
        "TENANT_SIG_NAMES": "ಪಾರ್ವತಮ್ಮ"
    }
    
    # Translate payload to Nudi ASCII
    translated_payload = {}
    for k, v in unicode_payload.items():
        translated_payload[k] = unicode_to_nudi(v)
        
    safe_uni = unicode_payload['OWNER_NAME'].encode('ascii', 'backslashreplace').decode('ascii')
    safe_nudi = translated_payload['OWNER_NAME'].encode('ascii', 'backslashreplace').decode('ascii')
    print(f"  Owner Unicode: {safe_uni} -> Nudi: {safe_nudi}")
    
    replace_placeholders(doc, translated_payload, font_name="Nudi Akshar-02")
    doc.save(output_docx)
    print(f"Saved generated document to {output_docx}")
    
    # Inspect generated paragraphs to verify replacement happened and font is set
    doc_new = docx.Document(output_docx)
    print("\nVerifying replaced runs:")
    found_owner = False
    for idx, p in enumerate(doc_new.paragraphs):
        # We check if owner placeholder is replaced (Nudi ASCII: µÀtÄäR¥Àà)
        if "µÀtÄäR¥Àà" in p.text:
            found_owner = True
            print(f"Paragraph {idx} contains owner name!")
            for r in p.runs:
                if "µÀtÄäR¥Àà" in r.text:
                    print(f"  Run text: {repr(r.text)} | Font Name: {r.font.name} | Size: {r.font.size}")
                    
    if found_owner:
        print("Success! Verification passed.")
    else:
        print("Warning: Replaced owner name not found in the output document.")

if __name__ == "__main__":
    test_kan_gen()
