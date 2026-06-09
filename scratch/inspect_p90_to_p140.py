import docx

doc = docx.Document(r"d:\rental pro\new\19-05-2026 - kan (2.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
for idx in range(90, min(140, len(doc.paragraphs))):
    p = doc.paragraphs[idx]
    text = p.text.strip()
    print(f"P {idx}: '{text}'")
