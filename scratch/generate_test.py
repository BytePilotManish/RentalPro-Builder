import sys
import os
import sqlite3
import json
import docx

# Add parent directory to sys.path
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
if parent_dir not in sys.path:
    sys.path.insert(0, parent_dir)

from app import replace_placeholders

def generate_local_test():
    conn = sqlite3.connect('database.db')
    cursor = conn.cursor()
    cursor.execute('SELECT id, title, agreement_data FROM agreements WHERE id = 7')
    row = cursor.fetchone()
    if not row:
        print("Agreement ID 7 not found.")
        return
        
    data = json.loads(row[2])
    print(f"Generating local test from data: {data}")
    
    doc = docx.Document("TEMPLATE.docx")
    replace_placeholders(doc, data)
    
    print("\n=== Paragraphs in Newly Generated Test DOCX ===")
    in_terms = False
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if "terms and conditions" in text.lower():
            in_terms = True
            print(f"P{idx}: {text}")
            continue
        if in_terms:
            if "IN WITNESSES THEREOF" in text:
                in_terms = False
                print(f"P{idx}: {text}")
                continue
            print(f"P{idx}: {text}")
            
    conn.close()

if __name__ == "__main__":
    generate_local_test()
