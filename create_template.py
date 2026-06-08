import os
import docx

def replace_text_in_doc(doc, replacements):
    # Process paragraphs
    for p in doc.paragraphs:
        text_before = p.text
        for placeholder, replacement in replacements.items():
            if placeholder in text_before:
                # Try replacing in runs first to preserve formatting
                for run in p.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, replacement)
                        # Ensure run inherits or explicitly gets 12pt Times New Roman
                        if run.font.size is None:
                            run.font.size = docx.shared.Pt(12)
                        if run.font.name is None:
                            run.font.name = 'Times New Roman'
                
                # If still present (due to placeholder split across runs), replace in paragraph.text
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, replacement)
                    # p.text assignment clears runs, format new runs created
                    for run in p.runs:
                        run.font.size = docx.shared.Pt(12)
                        run.font.name = 'Times New Roman'
                    
    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text_before = p.text
                    for placeholder, replacement in replacements.items():
                        if placeholder in text_before:
                            for run in p.runs:
                                if placeholder in run.text:
                                    run.text = run.text.replace(placeholder, replacement)
                                    if run.font.size is None:
                                        run.font.size = docx.shared.Pt(12)
                                    if run.font.name is None:
                                        run.font.name = 'Times New Roman'
                            if placeholder in p.text:
                                p.text = p.text.replace(placeholder, replacement)
                                for run in p.runs:
                                    run.font.size = docx.shared.Pt(12)
                                    run.font.name = 'Times New Roman'

def create_template():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    src_path = os.path.join(current_dir, "RENTAL AGREEMENT.docx")
    dest_path = os.path.join(current_dir, "TEMPLATE.docx")
    
    if not os.path.exists(src_path):
        print(f"Error: Source file '{src_path}' does not exist.")
        return
        
    doc = docx.Document(src_path)
    
    # Configure document Normal style to be 12 pt Times New Roman
    style = doc.styles['Normal']
    style.font.size = docx.shared.Pt(12)
    style.font.name = 'Times New Roman'
    
    # Define exact replacements to turn hardcoded text into clean template placeholders
    replacements = {
        "01ST Day of this April 2026": "{{AGREEMENT_DATE}}",
        "at Bangalore by and between:": "at {{AGREEMENT_PLACE}} by and between:",
        "Mr. MANOJ M & SANCHITHA C J": "{{OWNER_NAME}}",
        "S/O T Mahesh": "{{OWNER_PARENT}}",
        "Aged 45": "Aged {{OWNER_AGE}}",
        "No,99,100 C Near Sri Kalikamba Temple, ChowdeshwariNagar, , Laggere, Bengaluru- 560 058": "{{OWNER_ADDRESS}}",
        "Mr. RAJUGOWDA": "{{TENANT_NAME}}",
        "S/O Subbegowda": "{{TENANT_PARENT}}",
        "Aged 47": "Aged {{TENANT_AGE}}",
        "No. 18, 3rd Cross, Rajeev Gandhi Nagar, Laggere, Bengaluru-560 058": "{{TENANT_ADDRESS}}",
        "Commercial Purpose No.99 & 100C, Near Sri Kalikamba Temple Chowdeshwari Nagar, Laggere, Bengaluru- 560 058, Consists One RCC Roofed Shops, with rolling Shutter and electricity, Toilet and water facility": "Commercial Purpose {{PREMISES_ADDRESS}}, Consists {{PREMISES_DESCRIPTION}}",
        "his J S TRADERS": "his {{BUSINESS_NAME}}",
        "period of eleven months": "period of {{LEASE_PERIOD}}",
        "i.e., on 25/02/2027": "i.e., on {{LEASE_END_DATE}}",
        "Rs.1-10,500.00": "Rs.{{RENT_AMOUNT}}",
        "Rupees Ten Thousand Five Hundred only": "Rupees {{RENT_AMOUNT_WORDS}} only",
        "on 15th of every month": "on {{RENT_PAYMENT_DAY}} of every month",
        "pay 5% increase": "pay {{ESCALATION_RATE}} increase",
        "every 11 months, over": "every {{LEASE_PERIOD_NUM}} months, over",
        "further period of 11 months": "further period of {{LEASE_PERIOD_NUM}} months",
        "Rs 50,000/-": "Rs {{DEPOSIT_AMOUNT}}/-",
        "Rupees Fifty Thousand only": "Rupees {{DEPOSIT_AMOUNT_WORDS}} only",
        "way of cash as": "way of {{DEPOSIT_MODE}} as",
        "three months prior notice": "{{NOTICE_PERIOD}} prior notice",
        "(MANOJ M AND SANCHITHA C J)": "({{OWNER_SIG_NAMES}})",
        "(RAJUGOWDA)": "({{TENANT_SIG_NAMES}})"
    }
    
    replace_text_in_doc(doc, replacements)
    
    # Clean up non-unicode quotes
    for p in doc.paragraphs:
        for bad_char, good_char in [("\x93", '"'), ("\x94", '"'), ("\x91", "'"), ("\x92", "'"), ("\x96", "-")]:
            if bad_char in p.text:
                p.text = p.text.replace(bad_char, good_char)
                for run in p.runs:
                    run.font.size = docx.shared.Pt(12)
                    run.font.name = 'Times New Roman'
            
    # Find and replace the terms with {{AGREEMENT_CONDITIONS}} placeholder
    terms_indexes = []
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if (text.startswith("This RENTAL AGREEMENT is for a period") or
            text.startswith("The OWNER has agreed to let out") or
            text.startswith("The TENANT has agreed to pay the rental") or
            text.startswith("The TENANT has agreed to pay {{ESCALATION_RATE}}") or
            text.startswith("The TENANT has paid Security deposited") or
            text.startswith("The TENANT should use") or
            text.startswith("The TENANT should pay the Electricity") or
            text.startswith("The tenancy period may be renewed") or
            text.startswith("The OWNER and TENANT have agreed that {{NOTICE_PERIOD}}")):
            terms_indexes.append(idx)
            
    if terms_indexes:
        first_idx = terms_indexes[0]
        first_p = doc.paragraphs[first_idx]
        placeholder_p = first_p.insert_paragraph_before("{{AGREEMENT_CONDITIONS}}")
        placeholder_p.style = doc.styles['Normal']
        for run in placeholder_p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = docx.shared.Pt(12)
            
        for idx in sorted(terms_indexes, reverse=True):
            p_to_remove = doc.paragraphs[idx + 1]
            p_to_remove._element.getparent().remove(p_to_remove._element)
            
    doc.save(dest_path)
    print("TEMPLATE.docx created successfully!")

if __name__ == "__main__":
    create_template()
