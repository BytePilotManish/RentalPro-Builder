import os
import docx

def inspect_docx():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(current_dir, "RENTAL AGREEMENT.docx")
    
    if not os.path.exists(docx_path):
        print("Error: RENTAL AGREEMENT.docx not found.")
        return
        
    doc = docx.Document(docx_path)
    
    print(f"=== PARAGRAPHS ({len(doc.paragraphs)}) ===")
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            print(f"P{idx}: {text}")
            
    print(f"\n=== TABLES ({len(doc.tables)}) ===")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx}: {len(table.rows)} rows, {len(table.columns)} columns")
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"  R{r_idx}: {cells}")

if __name__ == "__main__":
    inspect_docx()
