import os
import docx

# Import the generation and placeholder replacement code from app.py
from app import replace_placeholders, AgreementData

def test_local_generation():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(current_dir, "TEMPLATE.docx")
    output_docx = os.path.join(current_dir, "output", "generated_test.docx")
    
    doc = docx.Document(template_path)
    
    # Standard template parameters (from test_generation.py)
    payload = {
        "AGREEMENT_DATE": "12th Day of June 2026",
        "AGREEMENT_PLACE": "Bengaluru",
        "OWNER_NAME": "Test Owner Name",
        "OWNER_PARENT": "S/O Test Owner Parent",
        "OWNER_AGE": "50",
        "OWNER_ADDRESS": "123 Owner Lane, Bengaluru",
        "TENANT_NAME": "Test Tenant Name",
        "TENANT_PARENT": "S/O Test Tenant Parent",
        "TENANT_AGE": "35",
        "TENANT_ADDRESS": "456 Tenant Road, Bengaluru",
        "PREMISES_ADDRESS": "Shop No 1, Main Bazaar, Bengaluru",
        "PREMISES_DESCRIPTION": "RCC Roofed Commercial Shop Premises",
        "BUSINESS_NAME": "SUPER TRADERS",
        "LEASE_PERIOD": "eleven months",
        "LEASE_PERIOD_NUM": "11",
        "LEASE_END_DATE": "11/05/2027",
        "RENT_AMOUNT": "12,000.00",
        "RENT_AMOUNT_WORDS": "Twelve Thousand only",
        "RENT_PAYMENT_DAY": "5th",
        "ESCALATION_RATE": "10%",
        "DEPOSIT_AMOUNT": "60,000",
        "DEPOSIT_AMOUNT_WORDS": "Sixty Thousand only",
        "DEPOSIT_MODE": "bank transfer",
        "NOTICE_PERIOD": "three months",
        "OWNER_SIG_NAMES": "Test Owner Name",
        "TENANT_SIG_NAMES": "Test Tenant Name"
    }
    
    replace_placeholders(doc, payload)
    
    # Ensure directory exists
    os.makedirs(os.path.dirname(output_docx), exist_ok=True)
    doc.save(output_docx)
    print("Saved generated_test.docx successfully.")
    
    # Now inspect fonts in generated_test.docx
    doc_g = docx.Document(output_docx)
    print(f"\n=== FONT SUMMARY FOR generated_test.docx ===")
    
    # Check default style
    try:
        print(f"Normal style font size: {doc_g.styles['Normal'].font.size}")
    except Exception as e:
        print(f"Error reading Normal style font size: {e}")
        
    font_sizes = {}
    for idx, p in enumerate(doc_g.paragraphs):
        for run in p.runs:
            sz = run.font.size
            if sz is not None:
                sz_pt = sz.pt
            else:
                sz_pt = "None"
            font_sizes[sz_pt] = font_sizes.get(sz_pt, 0) + 1
            
            # Print runs with size None or other than 12
            if sz_pt != 12.0 and run.text.strip():
                print(f"Paragraph {idx} Run: size={sz_pt}, text='{run.text.strip()[:30]}'")
                
    for table in doc_g.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        sz = run.font.size
                        if sz is not None:
                            sz_pt = sz.pt
                        else:
                            sz_pt = "None"
                        font_sizes[sz_pt] = font_sizes.get(sz_pt, 0) + 1
                        
    print("Run font sizes summary (in points):")
    for sz, count in font_sizes.items():
        print(f"  {sz} pt: {count} runs")

if __name__ == "__main__":
    test_local_generation()
