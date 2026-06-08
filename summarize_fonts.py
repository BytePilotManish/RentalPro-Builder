import os
import docx

def summarize_fonts(filename):
    current_dir = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(current_dir, filename)
    if not os.path.exists(path):
        print(f"File {filename} does not exist.")
        return
    
    doc = docx.Document(path)
    print(f"\n=== FONT SUMMARY FOR {filename} ===")
    
    # Check default style
    try:
        print(f"Normal style font size: {doc.styles['Normal'].font.size}")
    except Exception as e:
        print(f"Error reading Normal style font size: {e}")
        
    font_sizes = {}
    for p in doc.paragraphs:
        for run in p.runs:
            sz = run.font.size
            if sz is not None:
                sz_pt = sz.pt
            else:
                sz_pt = "None"
            font_sizes[sz_pt] = font_sizes.get(sz_pt, 0) + 1
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        sz = run.font.size
                        if sz is not None:
                            sz_pt = sz.pt
                        else:
                            sz_pt = "None"
                        font_sizes[sz_pt] = font_sizes.get(sz_pt, 0) + 1
                        
    print("Run font sizes (in points):")
    for sz, count in font_sizes.items():
        print(f"  {sz} pt: {count} runs")

if __name__ == "__main__":
    summarize_fonts("RENTAL AGREEMENT.docx")
    summarize_fonts("TEMPLATE.docx")
